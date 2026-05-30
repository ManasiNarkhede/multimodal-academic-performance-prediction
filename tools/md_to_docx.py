from pathlib import Path
from docx import Document
import markdown
import sys

md_path = Path('report/Final_Internship_Report_Expanded.md')
out_path = Path('report/Final_Internship_Report.docx')

if not md_path.exists():
    print(f"Markdown file not found: {md_path}")
    sys.exit(1)

text = md_path.read_text(encoding='utf-8')
# Convert Markdown to HTML then strip HTML tags? python-docx can't parse HTML.
# We'll convert headings and paragraphs roughly by parsing lines.

lines = text.splitlines()

doc = Document()

for line in lines:
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('#### '):
        doc.add_heading(line[5:].strip(), level=4)
    elif line.strip().startswith('- '):
        # simple bulleted list handling
        doc.add_paragraph(line.strip()[2:].strip(), style='List Bullet')
    elif line.strip() == '':
        doc.add_paragraph('')
    elif line.startswith('```'):
        # skip code fences markers
        continue
    else:
        doc.add_paragraph(line.strip())

# Save
out_path.parent.mkdir(parents=True, exist_ok=True)
doc.save(out_path)
print(f"Wrote {out_path}")
